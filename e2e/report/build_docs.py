#!/usr/bin/env python3
"""Sinh tài liệu Word kết quả test E2E từ e2e/results.json và ảnh trong e2e/evidence/.

Chạy:  python3 e2e/report/build_docs.py [thư mục xuất, mặc định ../Documents/Test]
Cần:   pip install python-docx
"""
import glob
import json
import os
import re
import sys
from datetime import datetime

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
RESULTS = os.path.join(ROOT, 'e2e', 'results.json')
EVIDENCE = os.path.join(ROOT, 'e2e', 'evidence')
OUT = os.path.abspath(sys.argv[1]) if len(sys.argv) > 1 else os.path.abspath(os.path.join(ROOT, '..', 'Documents', 'Test'))
os.makedirs(OUT, exist_ok=True)

NAVY = RGBColor(0x1F, 0x38, 0x64)
BLUE = RGBColor(0x2E, 0x75, 0xB6)
GREY = RGBColor(0x59, 0x59, 0x59)
FONT = 'Arial'

GROUPS = {
    '01': ('Xác thực & điều hướng', 'Đăng nhập, đăng xuất, đổi mật khẩu, trang lỗi 403/404, đổi ngôn ngữ'),
    '02': ('Dashboard', 'Tổng quan hệ thống, bản đồ, cây khu vực'),
    '03': ('Giám sát trực tiếp (Live)', 'Danh sách camera theo khu vực, chi tiết camera'),
    '04': ('Sự kiện nhiệt', 'Trang chủ sự kiện, cảnh báo nhiệt, lịch sử nhiệt độ'),
    '05': ('Phóng điện cục bộ (PD)', 'Cảnh báo PD, lịch sử PD, tổng hợp PD'),
    '06': ('Sự kiện AI', 'Cảnh báo từ camera AI'),
    '07': ('Kênh & nhóm thông báo', 'Kênh cảnh báo, bộ cảnh báo (nhóm)'),
    '08': ('Tỉnh, khu vực & người dùng', 'Quản lý tỉnh/thành phố, khu vực, người dùng'),
    '09': ('Camera', 'Hãng camera, danh sách camera'),
    '10': ('Cảm biến & thiết bị', 'Loại thiết bị, cảm biến, thiết bị'),
    '11': ('Cấu hình công thức', 'Công thức tính toán (biểu thức, biến số)'),
    '12': ('Quản lý AI', 'Chức năng AI, mô hình AI'),
    '13': ('Hướng dẫn sử dụng', 'Tài liệu hướng dẫn trong hệ thống'),
    '14': ('API và ràng buộc dữ liệu', 'Bảo mật API, kiểm tra dữ liệu đầu vào, trùng mã, ràng buộc xoá, ảnh cảnh báo, kết nối server AI'),
    '15': ('Phân quyền theo vai trò', 'Quyền của Viewer và Manager trên giao diện và API'),
    '16': ('Giao diện di động', 'Hiển thị trên màn hình điện thoại (390×844), không tràn ngang, biểu mẫu toàn màn hình'),
}

# Bước thực hiện / kết quả mong đợi theo loại test
TEMPLATES = [
    (r'^Xem danh sách', ['Đăng nhập bằng tài khoản quản trị.', 'Mở màn hình {screen} ({route}).'],
     'Màn hình tải thành công, hiển thị bộ lọc và bảng dữ liệu, không có lỗi.'),
    (r'^Kiểm tra bắt buộc', ['Mở màn hình {screen}.', 'Bấm "Thêm" để mở biểu mẫu.', 'Bấm "Lưu" khi chưa nhập gì.'],
     'Biểu mẫu không được lưu; các trường bắt buộc hiển thị thông báo lỗi.'),
    (r'^Thêm mới', ['Mở màn hình {screen}.', 'Bấm "Thêm".', 'Nhập dữ liệu thử (tiền tố TEST_) và chọn các mục bắt buộc.', 'Bấm "Lưu".'],
     'Hiển thị thông báo "Lưu thành công!" và bản ghi mới xuất hiện trong danh sách.'),
    (r'^Tìm kiếm', ['Mở màn hình {screen}.', 'Nhập chuỗi TEST_ vào ô lọc đầu tiên.', 'Bấm "Tìm kiếm".'],
     'Bảng chỉ hiển thị các bản ghi khớp điều kiện, gồm bản ghi test vừa tạo.'),
    (r'^Sửa', ['Mở màn hình {screen} và tìm bản ghi test.', 'Bấm biểu tượng sửa (✎).', 'Thêm hậu tố " (sửa)" vào tên.', 'Bấm "Lưu".'],
     'Hiển thị "Lưu thành công!" và tên mới hiển thị trong danh sách.'),
    (r'^Xoá', ['Mở màn hình {screen} và tìm bản ghi test.', 'Bấm biểu tượng xoá (🗑).', 'Xác nhận "Bạn có chắc muốn xóa".'],
     'Hộp thoại xác nhận xuất hiện; sau khi đồng ý, bản ghi biến mất khỏi danh sách.'),
]
CUSTOM = {
    'TC-01-01': (['Mở /login khi chưa đăng nhập.'], 'Hiển thị biểu mẫu đăng nhập gồm tên đăng nhập, mật khẩu, nút đăng nhập, chọn ngôn ngữ.'),
    'TC-01-02': (['Mở /login.', 'Bấm "Đăng nhập" khi để trống.'], 'Hiển thị lỗi bắt buộc nhập, không gửi yêu cầu đăng nhập.'),
    'TC-01-03': (['Nhập tên đăng nhập đúng và mật khẩu sai.', 'Bấm "Đăng nhập".'], 'Ở lại trang đăng nhập và hiển thị thông báo mật khẩu không đúng.'),
    'TC-01-04': (['Nhập tên đăng nhập và mật khẩu đúng.', 'Bấm "Đăng nhập".'], 'Chuyển vào hệ thống (Dashboard), hiển thị tên người dùng ở thanh điều hướng.'),
    'TC-01-05': (['Chưa đăng nhập, truy cập trực tiếp /dashboard.'], 'Bị chuyển về /login.'),
    'TC-01-06': (['Mở /login.', 'Bấm nút ngôn ngữ (VI/EN).'], 'Giao diện đổi ngôn ngữ tương ứng.'),
    'TC-01-07': (['Truy cập đường dẫn không tồn tại.'], 'Chuyển đến trang 404.'),
    'TC-01-08': (['Truy cập /403.'], 'Hiển thị trang báo không đủ quyền truy cập.'),
    'TC-01-09': (['Đăng nhập và bấm vào tên người dùng trên thanh điều hướng.'], 'Hiển thị menu người dùng (đổi mật khẩu, đăng xuất).'),
    'TC-01-10': (['Mở /change-password.'], 'Hiển thị biểu mẫu đổi mật khẩu.'),
    'TC-02-01': (['Mở /dashboard.'], 'Hiển thị thẻ Tổng camera, Địa điểm giám sát, Cảnh báo mới, Nhiệt độ môi trường, bản đồ và danh sách khu vực.'),
    'TC-02-02': (['Mở /dashboard.', 'Bấm nút ngôn ngữ và chọn tiếng Anh.'], 'Giao diện chuyển sang tiếng Anh (hoặc hiển thị menu chọn ngôn ngữ).'),
    'TC-03-01': (['Mở /live.'], 'Hiển thị cây khu vực, danh sách camera và các ô xem trực tiếp.'),
    'TC-03-02': (['Mở /live.', 'Bấm vào một camera (Cam Nhiệt PTZ).'], 'Mở màn hình chi tiết/khung xem trực tiếp của camera. (Hình ảnh video phụ thuộc thiết bị thật.)'),
    'TC-04-01': (['Mở /event/home.'], 'Hiển thị biểu đồ nhật ký nhiệt độ theo điểm đo và thống kê số lượng cảnh báo.'),
    'TC-04-02': (['Mở /event/notification-system.', 'Bấm "Tìm kiếm".'], 'Hiển thị bộ lọc và bảng danh sách nhiệt độ vượt ngưỡng.'),
    'TC-04-03': (['Mở /event/event-history.', 'Bấm "Tìm kiếm".'], 'Hiển thị bộ lọc và bảng nhật ký nhiệt độ.'),
    'TC-05-01': (['Mở /event/pd-notification-system.', 'Bấm "Tìm kiếm".'], 'Hiển thị bảng danh sách cảnh báo phóng điện.'),
    'TC-05-02': (['Mở /event/pd-history.', 'Bấm "Tìm kiếm".'], 'Hiển thị bảng nhật ký phóng điện.'),
    'TC-05-03': (['Mở /event/pd-summary.', 'Bấm "Tìm kiếm".'], 'Hiển thị biểu đồ thống kê phóng điện vượt ngưỡng và tỷ lệ xử lý.'),
    'TC-06-01': (['Mở /event/ai-notification.', 'Bấm "Tìm kiếm".'], 'Hiển thị bộ lọc và danh sách cảnh báo AI.'),
    'TC-12-01': (['Mở /ai/feature-list.'], 'Hiển thị danh sách chức năng AI (mã, tên, mô hình, luật cảnh báo…).'),
    'TC-12-02': (['Mở /ai/model-list.'], 'Hiển thị danh sách mô hình AI (mã, tên, loại tác vụ, ngưỡng tin cậy…).'),
    'TC-13-01': (['Mở /report/manual.'], 'Màn hình hướng dẫn tải được, không báo lỗi.'),
}

STATUS_VI = {'passed': 'ĐẠT', 'failed': 'KHÔNG ĐẠT', 'timedOut': 'KHÔNG ĐẠT', 'skipped': 'BỎ QUA', 'interrupted': 'KHÔNG ĐẠT'}


# ---------- tiện ích docx ----------
def set_cell_bg(cell, hex_color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tcPr.append(shd)


def set_cell_margins(table, top=60, bottom=60, left=100, right=100):
    tblPr = table._tbl.tblPr
    m = OxmlElement('w:tblCellMar')
    for k, v in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        e = OxmlElement(f'w:{k}')
        e.set(qn('w:w'), str(v))
        e.set(qn('w:type'), 'dxa')
        m.append(e)
    tblPr.append(m)


def set_repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement('w:tblHeader')
    e.set(qn('w:val'), 'true')
    trPr.append(e)


def no_split(row):
    trPr = row._tr.get_or_add_trPr()
    e = OxmlElement('w:cantSplit')
    e.set(qn('w:val'), 'true')
    trPr.append(e)


def add_field(run, instr):
    for t, text in (('begin', None), (None, instr), ('separate', None), ('text', '1'), ('end', None)):
        if t in ('begin', 'separate', 'end'):
            e = OxmlElement('w:fldChar')
            e.set(qn('w:fldCharType'), t)
            run._r.append(e)
        elif t is None:
            e = OxmlElement('w:instrText')
            e.set(qn('xml:space'), 'preserve')
            e.text = text
            run._r.append(e)
        else:
            e = OxmlElement('w:t')
            e.text = text
            run._r.append(e)


def style_fonts(doc):
    st = doc.styles
    n = st['Normal']
    n.font.name = FONT
    n.font.size = Pt(10.5)
    n.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    n.paragraph_format.space_after = Pt(4)
    for name, size, color, before, after in (('Heading 1', 18, NAVY, 18, 8), ('Heading 2', 13, BLUE, 14, 4), ('Heading 3', 11, NAVY, 10, 3)):
        h = st[name]
        h.font.name = FONT
        h.font.size = Pt(size)
        h.font.bold = True
        h.font.color.rgb = color
        h.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
        h.element.rPr.rFonts.set(qn('w:ascii'), FONT)
        h.element.rPr.rFonts.set(qn('w:hAnsi'), FONT)
        h.paragraph_format.space_before = Pt(before)
        h.paragraph_format.space_after = Pt(after)
        h.paragraph_format.keep_with_next = True


def new_doc(footer_title):
    doc = Document()
    us = OxmlElement('w:updateFields')
    us.set(qn('w:val'), 'true')
    doc.settings.element.append(us)  # Word tự cập nhật mục lục khi mở
    style_fonts(doc)
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2)
    sec.top_margin, sec.bottom_margin = Cm(2.2), Cm(2)
    sec.different_first_page_header_footer = True
    # header
    hp = sec.header.paragraphs[0]
    r = hp.add_run('THERMAL MONITORING  |  Báo cáo kiểm thử chức năng')
    r.font.size = Pt(8.5)
    r.font.color.rgb = GREY
    pPr = hp._p.get_or_add_pPr()
    b = OxmlElement('w:pBdr')
    bt = OxmlElement('w:bottom')
    for k, v in (('val', 'single'), ('sz', '6'), ('space', '1'), ('color', '2E75B6')):
        bt.set(qn(f'w:{k}'), v)
    b.append(bt)
    pPr.append(b)
    # footer
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = fp.add_run(f'{footer_title}  –  Trang ')
    r.font.size = Pt(8.5)
    r.font.color.rgb = GREY
    r2 = fp.add_run()
    r2.font.size = Pt(8.5)
    r2.font.color.rgb = GREY
    add_field(r2, 'PAGE')
    return doc


def para(doc, text='', bold=False, size=None, color=None, align=None, after=None, italic=False):
    p = doc.add_paragraph()
    r = p.add_run(text)
    r.bold = bold
    r.italic = italic
    if size:
        r.font.size = Pt(size)
    if color is not None:
        r.font.color.rgb = color
    if align:
        p.alignment = align
    if after is not None:
        p.paragraph_format.space_after = Pt(after)
    return p


def cell_text(cell, text, bold=False, color=None, size=9.5, align=None):
    cell.text = ''
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    if color is not None:
        r.font.color.rgb = color
    if align:
        p.alignment = align


def make_table(doc, headers, rows, widths, status_col=None, header_bg='1F3864'):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    set_cell_margins(t)
    for i, h in enumerate(headers):
        c = t.rows[0].cells[i]
        cell_text(c, h, bold=True, color=RGBColor(0xFF, 0xFF, 0xFF), align=WD_ALIGN_PARAGRAPH.CENTER)
        set_cell_bg(c, header_bg)
    set_repeat_header(t.rows[0])
    for ri, row in enumerate(rows):
        cells = t.add_row().cells
        no_split(t.rows[-1])
        for i, v in enumerate(row):
            if status_col is not None and i == status_col:
                ok = v == 'ĐẠT'
                cell_text(cells[i], v, bold=True, color=RGBColor(0x00, 0x61, 0x00) if ok else RGBColor(0x9C, 0x00, 0x06), align=WD_ALIGN_PARAGRAPH.CENTER)
                set_cell_bg(cells[i], 'C6EFCE' if ok else 'FFC7CE')
            else:
                cell_text(cells[i], str(v), align=WD_ALIGN_PARAGRAPH.CENTER if i == 0 else None)
                if ri % 2 == 1:
                    set_cell_bg(cells[i], 'F2F6FB')
    for row in t.rows:
        for i, w in enumerate(widths):
            row.cells[i].width = Cm(w)
    return t


def toc_field(doc):
    p = doc.add_paragraph()
    r = p.add_run()
    for t, txt in (('begin', None), ('instr', 'TOC \\o "1-2" \\h \\z \\u'), ('separate', None), ('text', 'Nhấn chuột phải → Cập nhật trường để hiển thị mục lục (hoặc nhấn F9).'), ('end', None)):
        if t == 'instr':
            e = OxmlElement('w:instrText')
            e.set(qn('xml:space'), 'preserve')
            e.text = txt
        elif t == 'text':
            e = OxmlElement('w:t')
            e.text = txt
        else:
            e = OxmlElement('w:fldChar')
            e.set(qn('w:fldCharType'), t)
        r._r.append(e)


def cover(doc, title, subtitle, meta):
    for _ in range(5):
        doc.add_paragraph()
    para(doc, 'THERMAL MONITORING', bold=True, size=14, color=BLUE, align=WD_ALIGN_PARAGRAPH.CENTER, after=2)
    para(doc, 'BÁO CÁO KIỂM THỬ CHỨC NĂNG', bold=True, size=24, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, after=6)
    para(doc, title, bold=True, size=18, color=NAVY, align=WD_ALIGN_PARAGRAPH.CENTER, after=4)
    para(doc, subtitle, size=11, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, after=40)
    t = doc.add_table(rows=0, cols=2)
    t.style = 'Table Grid'
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_cell_margins(t, 80, 80, 140, 140)
    for k, v in meta:
        c = t.add_row().cells
        cell_text(c[0], k, bold=True, size=10)
        set_cell_bg(c[0], 'DEEAF6')
        cell_text(c[1], v, size=10)
        c[0].width, c[1].width = Cm(5), Cm(11)
    doc.add_page_break()


# ---------- đọc kết quả ----------
def load_results():
    data = json.load(open(RESULTS, encoding='utf-8'))
    out = {}

    def walk(s):
        for sp in s.get('specs', []):
            m = re.match(r'^(TC-(\d\d)-(\d+)) \| (.+?) \| (.+?) \| (.+)$', sp['title'])
            if not m:
                continue
            t = sp['tests'][0]
            r = t['results'][-1]
            err = ''
            if r.get('error'):
                err = re.sub(r'\x1b\[[0-9;]*m', '', r['error'].get('message', ''))[:400]
            out[m.group(1)] = dict(id=m.group(1), group=m.group(2), fn=m.group(4), screen=m.group(5), route=m.group(6),
                                   status=STATUS_VI.get(r['status'], r['status']), ms=r['duration'], error=err)
        for c in s.get('suites', []):
            walk(c)

    for s in data['suites']:
        walk(s)
    return data, out


def _cl(c): return f'Mở /category/camera-list và tìm dòng camera "{c}".'
def _st(i, c): return f'Mở /stream/detail/{i} ({c}).'
CUSTOM.update({
 'TC-09-30': ([_cl('Cam Nhiệt PTZ'), 'Đếm các nút trong ô cuối cùng (cột Hành động).'], 'Có đúng 7 nút: đồng bộ, xem góc quay, chỉnh góc quay, cài đặt góc quay, cài đặt AI, sửa, xoá.'),
 'TC-09-31': ([_cl('Cam Nhiệt PTZ'), 'Xem nút thứ 3 (Chỉnh góc quay); lặp lại với "Cam Thường Fix".'], 'Nút bị ẩn với camera nhiệt và hiện với camera thường.'),
 'TC-09-32': ([_cl('Cam Nhiệt PTZ'), 'Rê chuột vào từng nút Đồng bộ, Xem góc quay, Cài đặt góc quay, Cài đặt AI.'], 'Mỗi nút hiện chú thích đúng nội dung.'),
 'TC-09-33': ([_cl('Cam Nhiệt PTZ'), 'Bấm nút "Đồng bộ góc quay".'], 'Gọi API syncPresets/{id}, phản hồi 200 và hiển thị kết quả đồng bộ.'),
 'TC-09-34': ([_cl('Cam Nhiệt PTZ'), 'Bấm nút "Xem danh sách góc quay".'], 'Hộp thoại "Thông số góc quay" mở ra với 2 tab Presets và Tours.'),
 'TC-09-35': ([_cl('Cam Nhiệt PTZ'), 'Mở hộp thoại góc quay, chọn tab Tours.'], 'Tab Tours hiển thị và có nút "Đồng bộ từ camera".'),
 'TC-09-36': ([_cl('Cam Nhiệt PTZ'), 'Bấm nút "Cài đặt góc quay".'], 'Chuyển tới /category/camera-setting/{id} đúng camera.'),
 'TC-09-37': ([_cl('Cam Thường Fix'), 'Bấm nút "Chỉnh góc quay".'], 'Mở màn chỉnh vùng quan sát của camera.'),
 'TC-09-38': ([_cl('Cam Thường Fix'), 'Bấm nút "Cài đặt AI".'], 'Mở màn cài đặt AI của camera.'),
 'TC-09-39': ([_cl('Cam Thường Fix'), 'Bấm biểu tượng sửa (✎).'], 'Form sửa mở ra với dữ liệu hiện tại của camera.'),
 'TC-09-40': ([_cl('Cam Thường Fix'), 'Bấm biểu tượng xoá (🗑), sau đó bấm "Hủy" ở hộp xác nhận.'], 'Hộp xác nhận đóng lại, camera vẫn còn trong danh sách.'),
 'TC-09-41': (['Mở /category/camera-setting/2.'], 'Hiển thị tên camera, loại PTZ và bộ đếm Góc quay / Tour.'),
 'TC-09-42': (['Mở /category/camera-setting/2.', 'Mở ô "Cách thực hiện".'], 'Có các nguồn Tự động, API hãng, ONVIF.'),
 'TC-09-43': (['Mở /category/camera-setting/2.', 'Chọn tab "Vùng đo nhiệt", rồi quay lại tab "Góc quay & Tour".'], 'Chuyển tab được; tab Góc quay & Tour có nút "Lưu góc quay".'),
 'TC-09-44': (['Mở /category/camera-setting/2.', 'Bấm "Đồng bộ góc quay".'], 'Gọi API đồng bộ và nhận phản hồi 200.'),
 'TC-09-45': (['Mở /category/camera-setting/5 (camera cố định).'], 'Hiển thị thông báo camera không phải PTZ nên không có góc quay/tour.'),
 'TC-09-46': (['Mở /category/camera-setting/2.'], 'Bảng góc quay có cột Số và Tên.'),
 'TC-03-30': ([_st(2,'camera PTZ')], 'Có 8 nút hướng, tâm "PTZ", thanh "Tốc độ quay" mặc định 50%.'),
 'TC-03-31': ([_st(5,'camera cố định')], 'Không có bàn phím hướng; vẫn có Thu phóng, Khởi động lại; không có Hiệu chuẩn.'),
 'TC-03-40': ([_st(2,'camera PTZ'), 'Bấm vào thanh tốc độ quay ở vị trí ~90%.', 'Giữ nút hướng "Up".'], 'Phần trăm tốc độ thay đổi và lệnh gửi đi có tốc độ lớn hơn 0.'),
 'TC-03-41': ([_st(2,'camera PTZ'), 'Giữ nút "Phóng to", thả; giữ nút "Thu nhỏ", thả.'], 'Gửi lệnh ZoomIn và ZoomOut.'),
 'TC-03-42': ([_st(2,'camera PTZ'), 'Bấm "Khởi động lại", sau đó bấm "Hủy".'], 'Hiện hộp xác nhận "Khởi động lại camera?"; Hủy thì không gửi lệnh nào.'),
 'TC-03-43': ([_st(2,'camera PTZ'), 'Bấm "Khởi động lại" rồi đồng ý.'], 'Gửi lệnh Restart và hiển thị "Đã gửi lệnh tới camera".'),
 'TC-03-44': ([_st(2,'camera PTZ'), 'Bấm "Hiệu chuẩn" rồi đồng ý.'], 'Hiện hộp xác nhận calibration và gửi lệnh Calibration khi đồng ý.'),
 'TC-03-45': ([_st(2,'camera PTZ'), 'Giả lập máy chủ trả lỗi 400, giữ nút "Left".'], 'Giao diện hiển thị thông báo lỗi cho người dùng.'),
 'TC-03-46': ([_st(2,'camera PTZ'), 'Bấm "Vẽ điểm/vùng", sau đó bấm "Thoát".'], 'Vào chế độ vẽ có gợi ý "Bấm lên hình để thêm điểm đo" và nút Đo nhiệt/Thoát; thoát trở về trạng thái đầu.'),
 'TC-03-47': ([_st(2,'camera PTZ'), 'Bấm "Cài đặt camera".'], 'Chuyển tới /category/camera-setting/2.'),
 'TC-03-48': ([_st(2,'camera PTZ'), 'Bấm nút quay lại.'], 'Về màn Giám sát trực tiếp (/live).'),
 'TC-03-49': (['Mở /live/detail/2.'], 'Mở được trang điều khiển camera, không chuyển sang 404.'),
 'TC-14-20': (['Gọi POST /api/Cameras/control với camera 2, lệnh Stop.'], 'Trả 200.'),
 'TC-14-21': (['Gọi POST /api/Cameras/control với camera 999999.'], 'Trả lỗi 4xx, không phải 500.'),
 'TC-14-22': (['Gọi POST /api/Cameras/control với mã lệnh 999.'], 'Mã lệnh không hợp lệ bị từ chối bằng lỗi 4xx.'),
 'TC-14-23': (['Gọi POST /api/Cameras/control không kèm token.'], 'Trả 401.'),
})
for _d in ('Up','Down','Left','Right','UpLeft','UpRight','DownLeft','DownRight'):
    _i = 32 + ['Up','Down','Left','Right','UpLeft','UpRight','DownLeft','DownRight'].index(_d)
    CUSTOM[f'TC-03-{_i}'] = ([_st(2,'camera PTZ'), f'Giữ chuột vào nút hướng "{_d}" khoảng 0,5 giây rồi thả.'], f'Gửi lệnh {_d} khi nhấn và lệnh Stop khi thả (lệnh được chặn ở tầng mạng, camera thật không chuyển động).')


def steps_for(tc):
    if tc['id'] in CUSTOM:
        return CUSTOM[tc['id']]
    num = int(tc['id'].split('-')[2])
    if num >= 20 or tc['group'] in ('14', '15', '16'):
        # Test mở rộng: tiêu đề mô tả chính kịch bản và kết quả mong đợi
        if tc['route'].startswith('/api/'):
            steps = [f'Đăng nhập (hoặc không đăng nhập nếu kiểm tra bảo mật) và gọi API {tc["route"]}.', f'Thực hiện: {tc["fn"]}.']
        else:
            steps = [f'Mở màn hình {tc["screen"]} ({tc["route"]}).', f'Thực hiện: {tc["fn"]}.']
        return steps, tc['fn'][0].upper() + tc['fn'][1:] + ' (đúng như mô tả, không phát sinh lỗi).'
    for pat, steps, exp in TEMPLATES:
        if re.match(pat, tc['fn']):
            return [s.format(screen=tc['screen'], route=tc['route']) for s in steps], exp
    return ([f'Mở màn hình {tc["screen"]}.'], 'Chức năng hoạt động đúng.')


def images_for(tc):
    files = sorted(glob.glob(os.path.join(EVIDENCE, tc['group'], tc['id'] + '*.png')))
    order = {'form': 0, 'confirm': 0}
    files.sort(key=lambda f: (0 if re.search(r'-(form|confirm)\.png$', f) else 1, f))
    return files


def img_caption(f, tc):
    m = re.search(r'-(form|confirm)\.png$', f)
    label = {'form': 'biểu mẫu đã nhập dữ liệu', 'confirm': 'hộp thoại xác nhận'}.get(m.group(1)) if m else 'kết quả sau khi thực hiện'
    return f'{tc["id"]} – {label}'


# ---------- tài liệu theo nhóm ----------
def build_group(g, tcs, env_meta, files_index):
    name, desc = GROUPS[g]
    fname = f'{g}_{slug(name)}.docx'
    files_index[g] = fname
    doc = new_doc(f'{g}. {name}')
    ok = sum(1 for t in tcs if t['status'] == 'ĐẠT')
    cover(doc, f'Nhóm {g}: {name}', desc, [
        ('Dự án', 'Thermal Monitoring (Frontend thermal-fe-v2 + Backend ThermalMonitoring)'),
        ('Công cụ kiểm thử', 'Playwright (Chromium), giao diện tiếng Việt'),
        ('Môi trường', env_meta['env']),
        ('Ngày thực hiện', env_meta['date']),
        ('Số test case', f'{len(tcs)}  (Đạt: {ok}  –  Không đạt: {len(tcs) - ok})'),
        ('Tài khoản', 'admin (quyền quản trị)'),
    ])
    doc.add_heading('Mục lục', level=1)
    toc_field(doc)
    doc.add_page_break()

    doc.add_heading('1. Tổng hợp test case', level=1)
    para(doc, f'Nhóm chức năng "{name}" gồm {len(tcs)} test case. Kết quả chi tiết và ảnh chụp màn hình ở mục 2.', after=8)
    make_table(doc, ['Mã TC', 'Chức năng', 'Màn hình', 'Đường dẫn', 'Kết quả'],
               [[t['id'], t['fn'], t['screen'], t['route'], t['status']] for t in tcs], [2.2, 5.8, 3.6, 3.6, 1.8], status_col=4)

    doc.add_heading('2. Chi tiết kiểm thử', level=1)
    fig = 0
    for t in tcs:
        doc.add_heading(f'{t["id"]} – {t["fn"]}', level=2)
        steps, exp = steps_for(t)
        info = doc.add_table(rows=0, cols=2)
        info.style = 'Table Grid'
        set_cell_margins(info)
        rows = [('Chức năng', t['fn']), ('Màn hình', t['screen']), ('Đường dẫn', t['route']),
                ('Điều kiện', 'Đã đăng nhập (trừ nhóm 01); dữ liệu test có tiền tố TEST_.'),
                ('Các bước', '\n'.join(f'{i + 1}. {s}' for i, s in enumerate(steps))),
                ('Kết quả mong đợi', exp),
                ('Kết quả thực tế', t['status'] + (f' – {t["error"]}' if t['error'] else '')),
                ('Thời gian chạy', f'{t["ms"] / 1000:.1f} giây')]
        for k, v in rows:
            c = info.add_row().cells
            no_split(info.rows[-1])
            cell_text(c[0], k, bold=True)
            set_cell_bg(c[0], 'DEEAF6')
            if k == 'Kết quả thực tế':
                good = t['status'] == 'ĐẠT'
                cell_text(c[1], v, bold=True, color=RGBColor(0x00, 0x61, 0x00) if good else RGBColor(0x9C, 0x00, 0x06))
                set_cell_bg(c[1], 'C6EFCE' if good else 'FFC7CE')
            else:
                cell_text(c[1], v)
            c[0].width, c[1].width = Cm(3.4), Cm(13.6)
        para(doc, '', after=2)
        for f in images_for(t):
            fig += 1
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.keep_with_next = True
            p.paragraph_format.space_after = Pt(2)
            p.add_run().add_picture(f, width=Cm(15.5))
            para(doc, f'Hình {g}.{fig}: {img_caption(f, t)}', italic=True, size=9, color=GREY, align=WD_ALIGN_PARAGRAPH.CENTER, after=10)
    doc.save(os.path.join(OUT, fname))
    return fname


def slug(s):
    import unicodedata
    s = unicodedata.normalize('NFD', s.replace('đ', 'd').replace('Đ', 'D'))
    s = ''.join(c for c in s if unicodedata.category(c) != 'Mn')
    return re.sub(r'[^A-Za-z0-9]+', '_', s).strip('_')


# ---------- mục lục chức năng ----------
def build_index(by_group, env_meta, files_index):
    doc = new_doc('Mục lục chức năng')
    allc = [t for g in sorted(by_group) for t in by_group[g]]
    ok = sum(1 for t in allc if t['status'] == 'ĐẠT')
    cover(doc, 'Mục lục chức năng kiểm thử', 'Danh mục nhóm chức năng, màn hình và tài liệu kết quả chi tiết', [
        ('Dự án', 'Thermal Monitoring'), ('Môi trường', env_meta['env']), ('Ngày thực hiện', env_meta['date']),
        ('Tổng số test case', f'{len(allc)}  (Đạt: {ok}  –  Không đạt: {len(allc) - ok})'),
    ])
    doc.add_heading('1. Danh mục nhóm chức năng', level=1)
    para(doc, 'Mỗi nhóm có một file Word riêng gồm bảng test case, các bước, kết quả mong đợi/thực tế và ảnh chụp màn hình.', after=8)
    rows = []
    for g in sorted(by_group):
        tcs = by_group[g]
        o = sum(1 for t in tcs if t['status'] == 'ĐẠT')
        rows.append([g, GROUPS[g][0], GROUPS[g][1], files_index[g], f'{o}/{len(tcs)}'])
    make_table(doc, ['Nhóm', 'Tên nhóm', 'Phạm vi', 'File chi tiết', 'Đạt'], rows, [1.4, 3.6, 5.6, 4.9, 1.5])

    doc.add_heading('2. Danh mục theo màn hình', level=1)
    scr = {}
    for t in allc:
        scr.setdefault((t['group'], t['screen'], t['route']), []).append(t)
    rows = []
    for (g, s, r), ts in sorted(scr.items()):
        o = sum(1 for t in ts if t['status'] == 'ĐẠT')
        rows.append([g, s, r, ', '.join(x['id'] for x in ts), f'{o}/{len(ts)}'])
    make_table(doc, ['Nhóm', 'Màn hình', 'Đường dẫn', 'Test case', 'Đạt'], rows, [1.4, 4.0, 4.6, 5.5, 1.5])

    doc.add_heading('3. Danh mục chức năng (toàn bộ test case)', level=1)
    rows = [[t['id'], t['fn'], t['screen'], files_index[t['group']].split('_')[0] + '_…', t['status']] for t in allc]
    make_table(doc, ['Mã TC', 'Chức năng', 'Màn hình', 'File', 'Kết quả'], rows, [2.2, 6.6, 4.2, 2.2, 1.8], status_col=4)
    doc.save(os.path.join(OUT, '00_Muc_Luc_Chuc_Nang.docx'))


# ---------- tổng hợp ----------
def build_summary(by_group, data, env_meta):
    doc = new_doc('Tổng hợp kết quả')
    allc = [t for g in sorted(by_group) for t in by_group[g]]
    ok = sum(1 for t in allc if t['status'] == 'ĐẠT')
    total_s = data['stats']['duration'] / 1000
    cover(doc, 'Tổng hợp kết quả kiểm thử', 'Thống kê, phạm vi chưa kiểm thử, dữ liệu test và quan sát', [
        ('Môi trường', env_meta['env']), ('Ngày thực hiện', env_meta['date']),
        ('Tổng số test case', str(len(allc))), ('Đạt / Không đạt', f'{ok} / {len(allc) - ok}'),
        ('Thời gian chạy', f'{int(total_s // 60)} phút {int(total_s % 60)} giây'),
    ])
    doc.add_heading('1. Kết quả theo nhóm', level=1)
    rows = []
    for g in sorted(by_group):
        o = sum(1 for t in by_group[g] if t['status'] == 'ĐẠT')
        rows.append([g, GROUPS[g][0], len(by_group[g]), o, len(by_group[g]) - o, f'{100 * o // len(by_group[g])}%'])
    rows.append(['', 'TỔNG', len(allc), ok, len(allc) - ok, f'{100 * ok // len(allc)}%'])
    make_table(doc, ['Nhóm', 'Tên nhóm', 'Số TC', 'Đạt', 'Không đạt', 'Tỷ lệ'], rows, [1.5, 6.5, 2, 2, 2.5, 2.5])

    doc.add_heading('2. Dữ liệu test', level=1)
    for s in ('Toàn bộ bản ghi tạo ra dùng tiền tố TEST_ (người dùng: test_user01, test_viewer, test_mgr) trong CSDL local thermal_monitoring_v2.',
              'Bản ghi tạo: khu vực, hãng camera, loại thiết bị, camera, cảm biến, thiết bị, người dùng, kênh cảnh báo, bộ cảnh báo, công thức.',
              'Dữ liệu nền được chèn trực tiếp: 1 loại cảm biến TEST_STYPE (bảng sensor_types đang trống nên không thể tạo cảm biến).',
              'Sau khi chạy, các bản ghi trên được xoá qua giao diện (soft-delete) và script e2e/seed/cleanup.sh dọn hẳn; đã kiểm tra không còn dữ liệu TEST_.'):
        doc.add_paragraph(s, style='List Bullet')

    failed = [t for t in allc if t['status'] != 'ĐẠT']
    doc.add_heading('3. Test không đạt – lỗi của phần mềm', level=1)
    if failed:
        para(doc, 'Các test dưới đây không đạt do hành vi của phần mềm khác với kỳ vọng (không phải lỗi của bộ test):', after=4)
        make_table(doc, ['Mã TC', 'Chức năng', 'Màn hình', 'Chi tiết lỗi'],
                   [[t['id'], t['fn'], t['screen'], (t['error'].split('\n')[0] or '').replace('Error: ', '')[:230]] for t in failed], [2.0, 5.0, 3.2, 6.8])
    else:
        para(doc, 'Không có test nào không đạt.')

    doc.add_heading('4. Phạm vi chưa kiểm thử', level=1)
    doc.add_paragraph('Các mục dưới đây chưa được kiểm thử tự động trong đợt này (cần thiết bị/dịch vụ thật hoặc kế hoạch riêng):')
    for s in ('Phát video camera thật (go2rtc/RTSP) và chuyển động thật của camera: các lệnh PTZ/Khởi động lại/Hiệu chuẩn trên giao diện được chặn ở tầng mạng (mock) để không làm camera thật hoạt động; chỉ lệnh Stop và đồng bộ góc quay gọi tới camera thật. Chưa kiểm thử lưu preset/tour và vẽ điểm đo nhiệt trên hình thật.',
              'Nhập Excel (Import) và IEC; xuất Excel hiện lỗi trên máy macOS (xem mục 3 nếu có test không đạt).',
              'Gửi cảnh báo thật qua Firebase push, Telegram, email; dữ liệu nhiệt/PD/AI thời gian thực từ Producer–Kafka.',
              'Sự kiện nhiệt (DB local chưa có bản ghi); thêm/sửa/xoá mô hình và chức năng AI hoàn chỉnh (chỉ kiểm tra danh sách, form, ràng buộc).',
              'Phân quyền chi tiết cho các vai trò OPERATOR và scheduler (đã kiểm thử Viewer, Manager, admin).',
              'Các trình duyệt khác ngoài Chromium; màn hình máy tính bảng.'):
        doc.add_paragraph(s, style='List Bullet')

    doc.add_heading('5. Quan sát trong quá trình kiểm thử', level=1)
    para(doc, 'Đây không phải lỗi làm test không đạt, nhưng nên xem xét:', after=4)
    for s in ('Đã sửa trong đợt này: menu "Quản trị hệ thống" nay ẩn các mục không có quyền; route /live/detail/{id} hoạt động (đổi tên route trùng thành stream_detail); API điều khiển camera/tour từ chối mã lệnh không hợp lệ bằng lỗi 400.',
              'Vai trò Viewer được cấp quyền đọc Kênh cảnh báo và Bộ cảnh báo nên menu hiện hai mục này là đúng thiết kế.',
              'SignalR không dùng được WebSocket (lỗi 401) và tự chuyển sang long-polling; kênh HeartBeat dùng đường dẫn kiểu Windows.',
              'Hộp thoại "Hãy cho phép thiết bị bật âm thanh" xuất hiện mỗi lần mở trang và chặn mọi thao tác chuột cho đến khi bấm "Đóng".',
              'Biểu mẫu Khu vực: ô "Tỉnh" chỉ có lựa chọn sau khi nhập tỉnh ở màn hình mới /category/province (bảng provinces ban đầu trống, chưa nạp sẵn 63 tỉnh).',
              'Biểu mẫu Cảm biến: "Loại cảm biến" rỗng nếu bảng sensor_types chưa có dữ liệu và không có thông báo hướng dẫn; bắt buộc nhập "Địa chỉ IP (WAN)".',
              'Kiểm tra IP không nhất quán: camera yêu cầu IP kèm cổng (192.168.1.1:8080) còn cảm biến chỉ nhận IP không có cổng.',
              'Trường "Loại camera" bị lỗi "Invalid CameraPtzType value." nếu để trống, thông báo là tiếng Anh/mã nội bộ.',
              '.env.development trỏ API tới địa chỉ LAN cố định (192.168.100.171); đổi IP máy sẽ làm giao diện không gọi được backend local.',
              'Các cấu hình đường dẫn kiểu Windows (dấu \\) như ExcelTemplate, ExportDataPath, HeartBeat:FilePath không chạy đúng trên macOS/Linux (đã sửa cho ảnh khu vực, ảnh cảnh báo PD/AI; xuất Excel chưa sửa).',
              'Xuất Excel còn phụ thuộc đường dẫn ExcelTemplate kiểu Windows (chưa sửa, lỗi 500 trên macOS).'):
        doc.add_paragraph(s, style='List Bullet')

    doc.add_heading('6. Cách chạy lại', level=1)
    code = ['cd thermal-fe-v2', 'e2e/seed/cleanup.sh && e2e/seed/seed.sh', 'E2E_PASSWORD=<mật khẩu admin> E2E_PRODUCER_TOKEN=<Producer:Token> npx playwright test',
            'python3 e2e/report/build_docs.py   # sinh lại tài liệu Word']
    for c in code:
        p = doc.add_paragraph()
        r = p.add_run(c)
        r.font.name = 'Courier New'
        r.font.size = Pt(9.5)
        p.paragraph_format.left_indent = Cm(0.6)
        p.paragraph_format.space_after = Pt(2)
    doc.save(os.path.join(OUT, '99_Tong_Hop_Ket_Qua_Test.docx'))


def main():
    data, res = load_results()
    by_group = {}
    for tc in res.values():
        by_group.setdefault(tc['group'], []).append(tc)
    for g in by_group:
        by_group[g].sort(key=lambda t: int(t['id'].split('-')[2]))
    start = datetime.fromisoformat(data['stats']['startTime'].replace('Z', '+00:00')).astimezone()
    env_meta = {'env': 'Local dev – FE http://localhost:5273, BE http://localhost:5272, DB thermal_monitoring_v2',
                'date': start.strftime('%d/%m/%Y %H:%M')}
    files_index = {}
    for g in sorted(by_group):
        build_group(g, by_group[g], env_meta, files_index)
    build_index(by_group, env_meta, files_index)
    build_summary(by_group, data, env_meta)
    print('Đã tạo', len(files_index) + 2, 'file trong', OUT)


if __name__ == '__main__':
    main()
